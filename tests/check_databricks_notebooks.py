from contextlib import redirect_stdout
from io import StringIO
from pathlib import Path
from shutil import copy2
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import MagicMock, patch
import ast
import nbformat
import pandas as pd
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
notebooks = root / "notebooks"
original = nbformat.read(root / "ABS Labour Force.ipynb", as_version=4)
original_cells = {c.id: c.source for c in original.cells if c.cell_type == "code"}
assert not (notebooks / "04 Publish Release.ipynb").exists()

# Validate the notebook structure and Python syntax.
for path in [root / "ABS Labour Force.ipynb", *notebooks.glob("*.ipynb")]:
    nb = nbformat.read(path, as_version=4)
    nbformat.validate(nb)
    for cell in nb.cells:
        if cell.cell_type == "code":
            ast.parse(cell.source)

with TemporaryDirectory() as folder:
    workspace = Path(folder)
    (workspace / "referencedatasets").mkdir()
    for name in ["Table 010.xlsx", "Interim Table 14.xlsx", "MLF1.xlsx"]:
        copy2(root / "referencedatasets" / name, workspace / "referencedatasets" / name)
    params = {"run_date": "2026-09-27", "job_run_id": "12345", "workspace_path": str(workspace)}

    for key, title, filename, parse_id, tidy_id in [
        ("table_010", "Table 010", "Table 010.xlsx", "c580e5a9", "839f822b"),
        ("table_014", "Table 014", "Interim Table 14.xlsx", "a8a775f8", "4e4cef74"),
        ("mlf1", "MLF1", "MLF1.xlsx", "e089167e", "16a10d4b"),
    ]:
        path = next(notebooks.glob(f"0* Prepare {title}.ipynb"))
        nb = nbformat.read(path, as_version=4)
        dbutils = MagicMock()
        dbutils.widgets.get.side_effect = params.__getitem__
        spark = MagicMock()
        response = MagicMock(content=(workspace / "referencedatasets" / filename).read_bytes())
        env = {"dbutils": dbutils, "spark": spark}

        # Keep each download at July while making its reference workbook end in June.
        if key == "mlf1":
            # Rebuild the temporary reference without Excel presentation pivots.
            with pd.ExcelFile(workspace / "referencedatasets" / filename) as book:
                sheet_names = book.sheet_names
                reference = pd.read_excel(book, sheet_name="Data 1", header=None)
            reference.loc[reference[0].eq(pd.Timestamp("2026-07-01")), 0] = pd.Timestamp("2026-06-01")
            with pd.ExcelWriter(workspace / "referencedatasets" / filename, engine="openpyxl") as writer:
                for sheet in sheet_names:
                    (reference if sheet == "Data 1" else pd.DataFrame()).to_excel(writer, sheet_name=sheet, header=False, index=False)
        else:
            book = load_workbook(workspace / "referencedatasets" / filename)
            book["Data1"].delete_rows(book["Data1"].max_row)
            book.save(workspace / "referencedatasets" / filename)
            book.close()

        # Run every cell with a reference workbook download and a mocked Spark destination.
        with patch("requests.get", return_value=response), redirect_stdout(StringIO()):
            for cell in nb.cells:
                if cell.cell_type == "code":
                    exec(compile(cell.source, str(path), "exec"), env)
        df = env["df"]
        assert env["csv_file"].exists()
        assert env["run"] == workspace / "outputs" / key / "2026-09-27_12345"
        assert env["release_month"] == "2026-07"
        assert "release_month" not in [call.args[0] for call in dbutils.widgets.get.call_args_list]
        manifest = pd.read_csv(env["run"] / "sources.csv")
        assert manifest.loc[0, "reference_period"] == "2026-07-01"
        assert manifest.loc[0, "output_file"] == env["csv_file"].name
        assert manifest.loc[0, "source_url"] == env["source_url"]
        assert not list(env["run"].glob("*.json"))
        dbutils.jobs.taskValues.set.assert_not_called()

        # Compare the cleaned data with the original notebook's processing.
        sources = pd.DataFrame([(title, filename, env["download_file"])], columns=["dataset", "baseline_file", "download_file"]).set_index("dataset")
        old_env = {**env, "sources": sources}
        with redirect_stdout(StringIO()):
            exec(compile(original_cells[parse_id], parse_id, "exec"), old_env)
            exec(compile(original_cells[tidy_id], tidy_id, "exec"), old_env)
        expected = old_env[key] if key != "mlf1" else old_env["mlf"].rename(columns={"region_code": "sa4_code", "region": "sa4_region"})
        pd.testing.assert_frame_equal(df, expected[df.columns])

        # Read the saved CSV with identifiers kept as text and compare its values.
        dtypes = {column: str for column in df.columns if column not in ["date", "value"]}
        saved = pd.read_csv(env["csv_file"], dtype=dtypes, parse_dates=["date"])
        pd.testing.assert_frame_equal(saved, df.reset_index(drop=True), check_dtype=False, check_exact=False, rtol=1e-12, atol=1e-12)

        # Confirm the final cell writes this DataFrame directly to its own Delta table.
        assert spark.createDataFrame.call_count == 1
        assert spark.createDataFrame.call_args.args[0] is df
        sdf = spark.createDataFrame.return_value
        sdf.__getitem__.return_value.cast.assert_called_once_with("date")
        sdf.withColumn.assert_called_once_with("date", sdf.__getitem__.return_value.cast.return_value)
        writer = sdf.withColumn.return_value.write
        writer.format.assert_called_once_with("delta")
        writer.format.return_value.mode.assert_called_once_with("overwrite")
        writer.format.return_value.mode.return_value.saveAsTable.assert_called_once_with(f"YOUR_CATALOG.YOUR_SCHEMA.{key}")
        spark.sql.assert_not_called()
        print(key, df.shape, "matches original processing; CSV and direct overwrite verified")

        # A failed Delta write must leave no completion record.
        failed = env["output"] / "2026-09-28_failed"
        failed.mkdir()
        failed_spark = MagicMock()
        failed_spark.createDataFrame.return_value.withColumn.return_value.write.format.return_value.mode.return_value.saveAsTable.side_effect = RuntimeError(
            "Test publication failure"
        )
        with TestCase().assertRaisesRegex(RuntimeError, "Test publication failure"):
            exec(nb.cells[14].source, {**env, "spark": failed_spark, "run": failed})
        assert not (failed / "sources.csv").exists()

        # Ignore incomplete runs and another table's later release; advance past completed July.
        other_key = "table_014" if key != "table_014" else "mlf1"
        other = workspace / "outputs" / other_key / "other_run"
        other.mkdir(parents=True, exist_ok=True)
        pd.DataFrame({"reference_period": ["2030-01-01"]}).to_csv(other / "sources.csv", index=False)
        with patch("pandas.read_excel", return_value=pd.DataFrame({"date": [pd.Timestamp("2026-06-01")]})), redirect_stdout(StringIO()):
            exec(nb.cells[4].source, env)
        assert env["release_month"] == "2026-08"

        # Select the latest completed release regardless of folder name or insertion order.
        later = env["output"] / "earlier_folder_name"
        later.mkdir()
        pd.DataFrame({"reference_period": ["2026-09-01"]}).to_csv(later / "sources.csv", index=False)
        with patch("pandas.read_excel", return_value=pd.DataFrame({"date": [pd.Timestamp("2026-06-01")]})), redirect_stdout(StringIO()):
            exec(nb.cells[4].source, env)
        assert env["release_month"] == "2026-10"

        # If only an older completion exists, the newer reference month determines the next release.
        (later / "sources.csv").unlink()
        pd.DataFrame({"reference_period": ["2026-05-01"]}).to_csv(env["run"] / "sources.csv", index=False)
        with patch("pandas.read_excel", return_value=pd.DataFrame({"date": [pd.Timestamp("2026-06-01")]})), redirect_stdout(StringIO()):
            exec(nb.cells[4].source, env)
        assert env["release_month"] == "2026-07"
        (other / "sources.csv").unlink()
        print(f"{title} automatic release selection and publication-failure handling verified")
print("Local checks passed. Actual Delta writes still need verification in Databricks.")
